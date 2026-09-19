"""Turn any attachment (txt / pdf / docx / xlsx) into a ParsedDoc.

A ParsedDoc carries the raw text, a list of (label, value) pairs, a guessed
document type, and — if the file can't be read — a readable error. Nothing
here decides what a field *means*; that's fields.py.
"""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import PurePath

from .fields import is_label

logging.getLogger("pypdf").setLevel(logging.ERROR)  # broken PDFs are handled below

DOC_TYPE_PATTERNS = [
    # order matters: "BILL OF LADING INSTRUCTION" is an SI, not a BL
    ("SI", r"shipping instruction|bl instruction|bill of lading instruction|b/l instruction|arahan penghantaran"),
    ("BL", r"bill of lading|\bdraft b/?l\b|提单"),
    ("COMMERCIAL_INVOICE", r"commercial invoice|\binvoice no\b"),
    ("PACKING_LIST", r"packing list"),
    ("CERTIFICATE_OF_ORIGIN", r"certificate of origin"),
]


@dataclass
class ParsedDoc:
    path: str
    fmt: str
    ok: bool = True
    error: str | None = None          # "empty" | "corrupt" | "no_text_layer" | "unsupported"
    doc_type: str = "UNKNOWN"
    text: str = ""
    pairs: list[tuple[str, str]] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {"path": self.path, "format": self.fmt, "ok": self.ok, "error": self.error,
                "doc_type": self.doc_type, "text": self.text[:6000]}


def detect_doc_type(text: str, filename: str = "") -> str:
    head = "\n".join(text.strip().splitlines()[:4]).lower()
    for dtype, pat in DOC_TYPE_PATTERNS:
        if re.search(pat, head):
            return dtype
    # fall back to the whole document, then to the file name
    low = text.lower()
    for dtype, pat in DOC_TYPE_PATTERNS:
        if re.search(pat, low):
            return dtype
    stem = PurePath(filename).stem.upper()
    if stem.endswith("_SI"):
        return "SI"
    if stem.endswith("_BL"):
        return "BL"
    return "UNKNOWN"


# ---------------------------------------------------------------- per format

def _pairs_from_lines(lines: list[str]) -> list[tuple[str, str]]:
    """Handles both `Label: value` (with indented continuation lines) and the
    PDF layout where a label sits alone on a line and the value follows."""
    pairs: list[tuple[str, str]] = []
    cur_label: str | None = None
    cur_vals: list[str] = []

    def flush():
        nonlocal cur_label, cur_vals
        if cur_label is not None:
            pairs.append((cur_label, "\n".join(v for v in cur_vals if v)))
        cur_label, cur_vals = None, []

    for raw in lines:
        line = raw.rstrip()
        if not line.strip() or set(line.strip()) <= set("=-_*"):
            continue
        stripped = line.strip()
        m = re.match(r"^([^:：]{2,70})[:：]\s*(.*)$", stripped)
        if m and is_label(m.group(1)) and not raw.startswith("  "):
            flush()
            cur_label, cur_vals = m.group(1).strip(), [m.group(2).strip()]
            continue
        if is_label(stripped) and len(stripped) < 70:
            flush()
            cur_label, cur_vals = stripped, []
            continue
        if cur_label is not None:
            cur_vals.append(stripped)
    flush()
    return pairs


def _parse_txt(data: bytes) -> tuple[str, list]:
    text = data.decode("utf-8", errors="replace")
    return text, _pairs_from_lines(text.splitlines())


def _parse_pdf(data: bytes) -> tuple[str, list]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    return text, _pairs_from_lines(text.splitlines())


def _parse_docx(data: bytes) -> tuple[str, list]:
    import docx
    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs]
    pairs = _pairs_from_lines(lines)
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[0] and cells[1]:
                pairs.append((cells[0], cells[1]))
            lines.append(" : ".join(cells))
    return "\n".join(lines), pairs


def _parse_xlsx(data: bytes) -> tuple[str, list]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    lines, pairs = [], []
    for ws in wb:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if not cells:
                continue
            lines.append(" : ".join(cells))
            if len(cells) >= 2:
                # xlsx packs multi-line values as "NAME | line2; line3"
                pairs.append((cells[0], cells[1].replace(" | ", "\n").replace("; ", "\n")))
    return "\n".join(lines), pairs


PARSERS = {"txt": _parse_txt, "pdf": _parse_pdf, "docx": _parse_docx, "xlsx": _parse_xlsx}


def parse_attachment(path: str, data: bytes) -> ParsedDoc:
    fmt = path.rsplit(".", 1)[-1].lower() if "." in path else ""
    doc = ParsedDoc(path=path, fmt=fmt)
    if not data:
        doc.ok, doc.error = False, "empty"
        return doc
    parser = PARSERS.get(fmt)
    if parser is None:
        doc.ok, doc.error = False, "unsupported"
        return doc
    try:
        doc.text, doc.pairs = parser(data)
    except Exception as exc:  # truncated / garbled files
        doc.ok, doc.error = False, "corrupt"
        doc.text = f"[could not open file: {type(exc).__name__}: {exc}]"
        return doc
    if not doc.text.strip():
        # e.g. an image-only scanned PDF: needs OCR / vision
        doc.ok, doc.error = False, "no_text_layer"
        return doc
    doc.doc_type = detect_doc_type(doc.text, path)
    return doc
